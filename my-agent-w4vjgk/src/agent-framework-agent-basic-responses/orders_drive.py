"""Lee el XLSX compartido de Drive. Sin API key, sin guardar copias en disco.

Solo funciona mientras el enlace ya permita lectura sin iniciar sesión.
No modifica permisos ni publica el archivo. No reutiliza datos antiguos si falla
la actualización. Cache breve en memoria para no descargar por cada mensaje.
"""
import io
import os
import re
import threading
import time
import warnings
from urllib.error import HTTPError, URLError
from urllib.request import urlopen
from openpyxl import load_workbook
from orders_lookup import buscar_pedido, validar_busqueda, normalizar_etapa
from orders_dates import COLUMNA_PREVISTO, cabeceras_con_previsto

CACHE_SECONDS = 60
MAX_BYTES = 10 * 1024 * 1024
_lock = threading.Lock()
_cache = None


def cargar_primera_hoja(file_id, header_row):
    """Descarga una vez por minuto y conserva solo los campos usados."""
    global _cache
    if not re.fullmatch(r"[A-Za-z0-9_-]{10,200}", file_id) or header_row < 1:
        raise ValueError("Configuración inválida")
    key = (file_id, header_row)
    with _lock:
        if _cache is not None and _cache[0] == key and time.monotonic() - _cache[1] < CACHE_SECONDS:
            return _cache[2]
        url = f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx"
        with urlopen(url, timeout=25) as response:
            raw = response.read(MAX_BYTES + 1)
        if len(raw) > MAX_BYTES or not raw.startswith(b"PK"):
            raise ValueError("El enlace no devolvió un XLSX válido")
        # read_only: las extensiones visuales no se modifican ni se vuelven a guardar.
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Unknown extension is not supported.*")
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            try:
                sheet = wb.worksheets[0]
                rows = sheet.iter_rows(min_row=header_row, values_only=True)
                headers = cabeceras_con_previsto(next(rows, ()))
                needed = {"NV", "Correo", "Estado Logistica", COLUMNA_PREVISTO}
                positions = [i for i, v in enumerate(headers)
                             if str(v).strip() in needed or normalizar_etapa(v) == "produccion"]
                compact_headers = [headers[i] for i in positions]
                compact_rows = [[row[i] if i < len(row) else None for i in positions] for row in rows]
                data = (compact_headers, compact_rows)
            finally:
                wb.close()
        _cache = (key, time.monotonic(), data)
        return data


def consultar_estado_pedido(nv: str = "", correo: str = "") -> dict:
    error = validar_busqueda(nv, correo)
    if error:
        return error
    file_id = os.getenv("ROCKERS_SHEET_ID", "").strip()
    if not file_id:
        return {"resultado": "configuracion_pendiente"}
    try:
        headers, rows = cargar_primera_hoja(file_id, int(os.getenv("ROCKERS_HEADER_ROW", "2")))
        return buscar_pedido(headers, rows, nv, correo)
    except HTTPError as exc:
        if exc.code in (401, 403, 404):
            return {"resultado": "archivo_no_accesible"}
        return {"resultado": "error_consulta"}
    except (URLError, TimeoutError, OSError):
        return {"resultado": "error_consulta"}
    except Exception:
        return {"resultado": "estructura_no_valida"}
