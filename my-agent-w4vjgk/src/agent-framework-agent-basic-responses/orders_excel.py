"""Consulta Excel local en modo lectura; solo expone tres campos."""
import os
from pathlib import Path
from datetime import datetime, date
from orders_lookup import buscar_pedido, validar_busqueda
from orders_dates import cabeceras_con_previsto

def consultar_estado_pedido(nv: str = "", correo: str = "") -> dict:
    if nv == "" and correo == "":
        return {"resultado": "nv_invalido"}
    error = validar_busqueda(nv, correo)
    if error:
        return error
    configured = os.getenv("ROCKERS_EXCEL_PATH", "").strip()
    if not configured:
        return {"resultado": "configuracion_pendiente"}
    path = Path(configured)
    if not path.is_absolute():
        path = Path(__file__).parent / path
    if not path.is_file():
        return {"resultado": "archivo_no_encontrado"}
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"resultado": "dependencias_pendientes"}
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        tab = os.getenv("ROCKERS_SHEET_TAB", "").strip() or workbook.sheetnames[0]
        if tab not in workbook.sheetnames:
            return {"resultado": "estructura_no_valida"}
        header_row = int(os.getenv("ROCKERS_HEADER_ROW", "2"))
        if header_row < 1:
            return {"resultado": "estructura_no_valida"}
        rows = workbook[tab].iter_rows(min_row=header_row, values_only=True)
        headers = cabeceras_con_previsto(next(rows, ()))
        def convert(value):
            if isinstance(value, (datetime, date)):
                return value.strftime("%d-%m-%Y")
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return value
        return buscar_pedido(headers, ([convert(v) for v in row] for row in rows), nv, correo)
    except ValueError:
        return {"resultado": "estructura_no_valida"}
    except Exception:
        return {"resultado": "error_lectura"}
    finally:
        if workbook is not None:
            workbook.close()
